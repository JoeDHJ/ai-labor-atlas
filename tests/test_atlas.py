import csv
import io
import json
import tempfile
import unittest
import zipfile
from contextlib import redirect_stdout
from pathlib import Path

from ai_labor_atlas.cli import main as cli_main
from ai_labor_atlas.dashboard import render
from ai_labor_atlas.demo import TASK_FIELDS, FIELDS, demo_rows, demo_tasks
from ai_labor_atlas.distance import OccupationBridge
from ai_labor_atlas.metrics import group_by_major_soc, summarize, summarize_tasks
from ai_labor_atlas.occupation_context import suggest_occupations
from ai_labor_atlas.pipeline import (
    _load_oews,
    _load_projections,
    _load_tasks,
    build_dataset,
)
from ai_labor_atlas.reviews import (
    REVIEW_DISCLOSURE,
    load_reviews,
    normalize_review,
    summarize_reviews,
    validate_review_file,
)

try:
    from openpyxl import Workbook
except ImportError:  # pragma: no cover - optional Excel dependency
    Workbook = None


class AtlasTests(unittest.TestCase):
    def test_demo_has_contract_and_summary(self):
        rows = demo_rows()
        self.assertEqual(len(rows), 8)
        self.assertEqual(set(rows[0]), set(FIELDS))
        result = summarize(
            [{key: str(value) for key, value in row.items()} for row in rows]
        )
        self.assertEqual(result["rows"], 8)
        self.assertGreater(result["exposure_coverage"], 0.99)

        tasks = demo_tasks()
        self.assertEqual(set(tasks[0]), set(TASK_FIELDS))
        task_summary = summarize_tasks(
            [{key: str(value) for key, value in row.items()} for row in rows],
            [{key: str(value) for key, value in row.items()} for row in tasks],
        )
        self.assertEqual(task_summary["task_rows"], 16)
        self.assertEqual(task_summary["task_occupation_coverage"], 1.0)

    def test_demo_build_is_deterministic(self):
        with tempfile.TemporaryDirectory() as temp:
            processed = Path(temp) / "processed"
            manifest = build_dataset(Path(temp) / "raw", processed, demo=True)
            self.assertEqual(manifest["row_count"], 8)
            self.assertEqual(manifest["task_row_count"], 16)
            with (processed / "occupations.csv").open(
                newline="", encoding="utf-8"
            ) as handle:
                self.assertEqual(len(list(csv.DictReader(handle))), 8)
            with (processed / "tasks.csv").open(newline="", encoding="utf-8") as handle:
                self.assertEqual(len(list(csv.DictReader(handle))), 16)

    def test_task_statements_are_loaded_with_provenance(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "Task Statements.txt").write_text(
                "\t".join(
                    [
                        "O*NET-SOC Code",
                        "Task ID",
                        "Task",
                        "Task Type",
                        "Incumbents Responding",
                        "Date",
                        "Domain Source",
                    ]
                )
                + "\n"
                + "\t".join(
                    [
                        "15-1252.00",
                        "8823",
                        "Build software systems.",
                        "Core",
                        "95",
                        "08/2023",
                        "Incumbent",
                    ]
                )
                + "\n",
                encoding="utf-8",
            )
            tasks = _load_tasks(root)
            self.assertEqual(len(tasks), 1)
            self.assertEqual(tasks[0]["task_id"], "8823")
            self.assertEqual(tasks[0]["source_file"], "Task Statements.txt")

    def test_dashboard_includes_task_filter_and_task_payload(self):
        rows = [{key: str(value) for key, value in row.items()} for row in demo_rows()]
        tasks = [
            {key: str(value) for key, value in row.items()} for row in demo_tasks()
        ]
        page = render(rows, summarize(rows), [], tasks)
        self.assertIn('id="task-filter"', page)
        self.assertIn('id="task-type-filter"', page)
        self.assertIn("Core tasks", page)
        self.assertIn("tasks_by_onet", page)
        self.assertIn("Design, develop, and test software applications.", page)
        self.assertIn('id="bridge-select"', page)
        self.assertIn("Career bridge", page)

    def test_worker_review_contract_preserves_source_and_does_not_score_reviews(self):
        review = normalize_review(
            {
                "review_id": "review-1",
                "onet_soc_code": "15-2051.00",
                "source": "reddit",
                "source_url": "https://www.reddit.com/r/datascience/",
                "review_scope": "employer_role",
                "review_date": "2025-04-03",
                "job_title": "Data Scientist",
                "topics": ["pay_benefits", "work_environment"],
                "rating": 3,
                "excerpt": "The work was interesting, but the team was understaffed.",
                "private_moderation_note": "must never be public",
            }
        )
        context = summarize_reviews([review], "15-2051.00")
        self.assertEqual(context["review_count"], 1)
        self.assertEqual(context["reviews"][0]["source"], "reddit")
        self.assertNotIn("private_moderation_note", context["reviews"][0])
        self.assertNotIn("overall_rating", context)
        self.assertIn("biased", context["disclosure"])
        self.assertEqual(context["total_review_count"], 1)
        self.assertFalse(context["is_truncated"])

    def test_review_loader_missing_file_is_empty_and_invalid_source_is_rejected(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "reviews.json"
            self.assertEqual(load_reviews(path), [])
            path.write_text(
                '{"reviews": [{"review_id": "r", "onet_soc_code": "15-2051.00", "excerpt": "A comment", "source": "unknown"}]}',
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "unsupported review source"):
                load_reviews(path)

    def test_review_loader_rejects_duplicate_ids_and_non_onet_codes(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "reviews.json"
            path.write_text(
                '{"reviews": [{"review_id": "r", "onet_soc_code": "15-2051.00", "excerpt": "One", "source": "reddit"}, {"review_id": "r", "onet_soc_code": "15-2051.00", "excerpt": "Two", "source": "reddit"}]}',
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "review_id values must be unique"):
                load_reviews(path)
            path.write_text(
                '{"reviews": [{"review_id": "r-2", "onet_soc_code": "15-2051", "excerpt": "One", "source": "reddit"}]}',
                encoding="utf-8",
            )
            with self.assertRaisesRegex(ValueError, "O\\*NET format"):
                load_reviews(path)

    def test_review_validation_reports_all_errors_and_coverage(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "reviews.json"
            path.write_text(
                json.dumps(
                    [
                        {
                            "review_id": "r-1",
                            "onet_soc_code": "15-2051.00",
                            "source": "reddit",
                            "topics": ["work_environment"],
                            "excerpt": "The team was supportive.",
                        },
                        {
                            "review_id": "r-1",
                            "onet_soc_code": "15-2051.00",
                            "source": "reddit",
                            "excerpt": "The same source ID is duplicated.",
                        },
                        {
                            "review_id": "r-3",
                            "onet_soc_code": "15-2051",
                            "source": "indeed",
                            "excerpt": "The occupation code is incomplete.",
                        },
                    ]
                ),
                encoding="utf-8",
            )
            report = validate_review_file(path)
        self.assertFalse(report["valid"])
        self.assertEqual(report["row_count"], 3)
        self.assertEqual(report["normalized_row_count"], 2)
        self.assertEqual(report["invalid_row_count"], 3)
        self.assertEqual(report["duplicate_review_ids"], ["r-1"])
        self.assertEqual(report["occupation_codes"], ["15-2051.00"])
        self.assertEqual(report["source_counts"], {"reddit": 2})
        self.assertEqual(report["error_count"], 3)
        self.assertTrue(any(item["field"] == "onet_soc_code" for item in report["errors"]))
        self.assertNotIn("The team was supportive", json.dumps(report))

    def test_review_validation_reports_malformed_json_without_traceback_data(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "reviews.json"
            path.write_text("{not valid json", encoding="utf-8")
            report = validate_review_file(path)
        self.assertFalse(report["valid"])
        self.assertEqual(report["errors"][0]["field"], "file")
        self.assertIn("invalid review JSON", report["errors"][0]["message"])

    def test_validate_reviews_cli_returns_json_and_nonzero_for_invalid_file(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "reviews.json"
            path.write_text(
                json.dumps(
                    [
                        {
                            "review_id": "r-1",
                            "onet_soc_code": "15-2051",
                            "excerpt": "Needs correction.",
                        }
                    ]
                ),
                encoding="utf-8",
            )
            output = io.StringIO()
            with redirect_stdout(output):
                code = cli_main(["validate-reviews", "--input", str(path)])
        self.assertEqual(code, 2)
        report = json.loads(output.getvalue())
        self.assertFalse(report["valid"])
        self.assertEqual(report["errors"][0]["row"], 1)

    def test_validate_reviews_cli_accepts_valid_file(self):
        with tempfile.TemporaryDirectory() as temp:
            path = Path(temp) / "reviews.json"
            path.write_text(
                json.dumps(
                    [
                        {
                            "review_id": "r-1",
                            "onet_soc_code": "15-2051.00",
                            "source": "user_submitted",
                            "topics": ["growth"],
                            "excerpt": "The work offers room to learn.",
                        }
                    ]
                ),
                encoding="utf-8",
            )
            output = io.StringIO()
            with redirect_stdout(output):
                code = cli_main(["validate-reviews", "--input", str(path)])
        self.assertEqual(code, 0)
        report = json.loads(output.getvalue())
        self.assertTrue(report["valid"])
        self.assertEqual(report["occupation_count"], 1)

    def test_dashboard_includes_transparent_worker_review_slot(self):
        rows = [{key: str(value) for key, value in row.items()} for row in demo_rows()]
        page = render(rows, summarize(rows), [], demo_tasks())
        self.assertIn("What workers say", page)
        self.assertIn("About these reviews", page)
        self.assertIn('id="worker-review-source-filter"', page)
        self.assertIn('id="worker-review-topic-filter"', page)
        self.assertIn(REVIEW_DISCLOSURE, page)

    def test_occupation_suggestions_require_confirmation(self):
        rows = [{key: str(value) for key, value in row.items()} for row in demo_rows()]
        result = suggest_occupations(rows, "Software Developer")
        self.assertTrue(result["candidates"])
        self.assertTrue(result["requires_confirmation"])
        self.assertTrue(result["candidates"][0]["requires_confirmation"])

    def test_occupation_suggestions_reject_generic_or_incompatible_title_overlap(self):
        rows = [
            {key: str(value) for key, value in row.items()}
            for row in demo_rows()
        ]
        analyst = suggest_occupations(rows, "Data Analyst")
        self.assertEqual(analyst["candidates"], [])
        product = suggest_occupations(rows, "Product Manager")
        self.assertNotIn(
            "Demonstrators and Product Promoters",
            {item["title"] for item in product["candidates"]},
        )
        machine_learning = suggest_occupations(rows, "Machine Learning Engineer")
        self.assertEqual(machine_learning["candidates"], [])

    def test_occupation_suggestions_normalize_plural_title_phrases(self):
        rows = [{key: str(value) for key, value in row.items()} for row in demo_rows()]
        result = suggest_occupations(rows, "Software Developer")
        self.assertEqual(result["candidates"][0]["title"], "Software Developers")
        self.assertEqual(result["candidates"][0]["match_score"], 1.0)
        self.assertEqual(result["candidates"][0]["basis"], ["title_phrase_match"])

    def test_review_summary_exposes_source_and_topic_labels(self):
        context = summarize_reviews([], "15-2051.00")
        self.assertEqual(context["source_labels"]["reddit"], "Reddit")
        self.assertEqual(context["topic_labels"]["work_environment"], "Work environment")

    def test_review_summary_reports_when_display_limit_truncates_comments(self):
        reviews = [
            normalize_review(
                {
                    "review_id": f"review-{index}",
                    "onet_soc_code": "15-2051.00",
                    "source": "user_submitted",
                    "review_date": f"2025-01-{index:02d}",
                    "excerpt": f"Comment {index}",
                }
            )
            for index in range(1, 4)
        ]
        context = summarize_reviews(reviews, "15-2051.00", limit=2)
        self.assertEqual(context["review_count"], 2)
        self.assertEqual(context["total_review_count"], 3)
        self.assertTrue(context["is_truncated"])

    def test_occupation_bridge_keeps_distance_evidence_separate(self):
        with tempfile.TemporaryDirectory() as temp:
            raw = Path(temp)
            header = "O*NET-SOC Code\tElement ID\tElement Name\tScale ID\tData Value\n"
            for filename, element, name in [
                ("Essential Skills.txt", "2.A.1.a", "Reading Comprehension"),
                ("Transferable Skills.txt", "2.B.1.a", "Social Perceptiveness"),
                ("Knowledge.txt", "2.C.1.a", "Administration and Management"),
                ("Abilities.txt", "1.A.1.a", "Oral Comprehension"),
                ("Work Activities.txt", "4.A.1.a.1", "Getting Information"),
            ]:
                (raw / filename).write_text(
                    header
                    + f"15-1252.00\t{element}\t{name}\tIM\t4.0\n"
                    + f"13-2011.00\t{element}\t{name}\tIM\t3.0\n",
                    encoding="utf-8",
                )
            (raw / "Software Skills.txt").write_text(
                "O*NET-SOC Code\tWorkplace Example\tElement ID\tElement Name\tHot Technology\tIn Demand\n"
                "15-1252.00\tPython\t2.E.1.a\tProgramming\tY\tY\n"
                "13-2011.00\tMicrosoft Excel\t2.E.1.a\tSpreadsheet\tY\tY\n",
                encoding="utf-8",
            )
            rows = [
                {key: str(value) for key, value in row.items()}
                for row in demo_rows()[:2]
            ]
            bridge = OccupationBridge(raw, rows, demo_tasks()[:4])
            result = bridge.bridge(rows[0]["onet_soc_code"])
            self.assertTrue(result["available"])
            self.assertEqual(result["source"]["profile_source"], "exact")
            self.assertTrue(result["candidates"])
            candidate = result["candidates"][0]
            self.assertIn("structured_distance", candidate)
            self.assertIn("software_overlap", candidate)
            self.assertIn("shared_task_evidence", candidate)

    def test_occupation_bridge_is_explicit_when_structured_data_is_missing(self):
        rows = [{key: str(value) for key, value in row.items()} for row in demo_rows()]
        bridge = OccupationBridge(Path("Z:\\missing-onet"), rows, demo_tasks())
        result = bridge.bridge(rows[0]["onet_soc_code"])
        self.assertFalse(result["available"])
        self.assertEqual(result["error"], "structured_onet_profiles_unavailable")

    def test_grouping(self):
        rows = [{key: str(value) for key, value in row.items()} for row in demo_rows()]
        groups = group_by_major_soc(rows)
        self.assertTrue(any(group["soc_major"] == "15" for group in groups))

    @unittest.skipUnless(Workbook, "openpyxl is required for Excel parser tests")
    def test_bls_excel_inputs_are_parsed_with_explicit_units(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            oews_zip = root / "bls_oews_may_2025.zip"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["OCC_CODE", "TOT_EMP", "A_MEDIAN"])
            sheet.append(["15-1252", 1000, 50000])
            content = io.BytesIO()
            workbook.save(content)
            with zipfile.ZipFile(oews_zip, "w") as archive:
                archive.writestr("national.xlsx", content.getvalue())

            projection_path = root / "bls_projections_2024_2034.xlsx"
            projection_book = Workbook()
            projection_book.remove(projection_book.active)
            projection_sheet = projection_book.create_sheet("Table 1.2")
            projection_sheet.append(["Table 1.2 title"])
            projection_sheet.append(
                [
                    "2024 National Employment Matrix title",
                    "2024 National Employment Matrix code",
                    "Employment, 2024",
                    "Employment, 2034",
                    "Employment change, percent, 2024-34",
                    "Occupational openings, 2024-34 annual average",
                ]
            )
            projection_sheet.append(
                ["Software Developers", "15-1252", 10.0, 12.0, 20.0, 1.2]
            )
            projection_book.save(projection_path)

            oews = _load_oews(root)
            projections = _load_projections(root)
            self.assertEqual(oews["15-1252"]["employment_2024"], 1000)
            self.assertEqual(projections["15-1252"]["annual_openings_2024_2034"], 1200)


if __name__ == "__main__":
    unittest.main()

from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from pathlib import Path
from typing import Iterable

from .demo import FIELDS, TASK_FIELDS, demo_rows, demo_tasks
from .io import read_csv, sha256, write_csv, write_json
from .occupation_context import load_alias_registry, validate_alias_registry


def _value(row: dict[str, object], *names: str) -> str:
    def clean(value: object) -> str:
        return re.sub(r"[^a-z0-9]+", " ", str(value).strip().lower()).strip()

    lowered = {clean(key): value for key, value in row.items()}
    for name in names:
        key = clean(name)
        if key in lowered and lowered[key] not in (None, ""):
            return str(lowered[key]).strip()
    return ""


def _number(value: object, *, non_negative: bool = False) -> float | None:
    if value in (None, "", "*", "**", "#", "–", "—"):
        return None
    cleaned = re.sub(r"[^0-9.\-]", "", str(value))
    try:
        parsed = float(cleaned) if cleaned else None
        if parsed is None or not math.isfinite(parsed):
            return None
        return None if non_negative and parsed < 0 else parsed
    except ValueError:
        return None


def _read_tsv(path: Path) -> list[dict[str, str]]:
    with path.open(newline="", encoding="utf-8-sig", errors="replace") as handle:
        return list(csv.DictReader(handle, delimiter="\t"))


def _find_file(root: Path, names: Iterable[str]) -> Path | None:
    wanted = {name.lower() for name in names}
    for path in root.rglob("*"):
        if path.is_file() and path.name.lower() in wanted:
            return path
    return None


def _extract_onet_zip(zip_path: Path, temp_dir: Path) -> Path:
    temp_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(zip_path) as archive:
        archive.extractall(temp_dir)
    return temp_dir


def _load_onet(raw_dir: Path) -> list[dict[str, object]]:
    occupation_file = _find_file(
        raw_dir, ["Occupation Data.txt", "occupation data.txt"]
    )
    if not occupation_file:
        zip_path = raw_dir / "onet_30_3_text.zip"
        if zip_path.exists():
            _extract_onet_zip(zip_path, raw_dir / "onet_30_3_text")
            occupation_file = _find_file(
                raw_dir / "onet_30_3_text",
                ["Occupation Data.txt", "occupation data.txt"],
            )
    if not occupation_file:
        raise FileNotFoundError(
            "O*NET occupation data not found. Put Occupation Data.txt or onet_30_3_text.zip in data/raw."
        )
    rows = []
    for row in _read_tsv(occupation_file):
        code = _value(row, "O*NET-SOC Code", "O*NET SOC Code")
        if not code or code.lower().startswith("code"):
            continue
        rows.append(
            {
                "onet_soc_code": code,
                "title": _value(row, "Title"),
                "description": _value(row, "Description"),
            }
        )
    return rows


def _load_tasks(raw_dir: Path) -> list[dict[str, str]]:
    task_file = _find_file(raw_dir, ["Task Statements.txt", "task statements.txt"])
    if not task_file:
        zip_path = raw_dir / "onet_30_3_text.zip"
        if zip_path.exists():
            _extract_onet_zip(zip_path, raw_dir / "onet_30_3_text")
            task_file = _find_file(
                raw_dir / "onet_30_3_text",
                ["Task Statements.txt", "task statements.txt"],
            )
    if not task_file:
        raise FileNotFoundError(
            "O*NET task data not found. Put Task Statements.txt or onet_30_3_text.zip in data/raw."
        )
    rows = []
    for row in _read_tsv(task_file):
        code = _value(row, "O*NET-SOC Code", "O*NET SOC Code")
        task_id = _value(row, "Task ID", "task id")
        statement = _value(row, "Task", "Task Statement", "task statement")
        if not code or not task_id or not statement:
            continue
        rows.append(
            {
                "onet_soc_code": code,
                "task_id": task_id,
                "task_statement": statement,
                "task_type": _value(row, "Task Type", "task type"),
                "incumbents_responding": _value(
                    row, "Incumbents Responding", "incumbents responding"
                ),
                "task_date": _value(row, "Date", "task date"),
                "domain_source": _value(row, "Domain Source", "domain source"),
                "onet_version": "30.3",
                "source_file": task_file.name,
                "task_quality_flags": "",
            }
        )
    return rows


def _load_crosswalk(raw_dir: Path) -> dict[str, list[str]]:
    workbook = raw_dir / "onet_soc_2019_to_2018.xlsx"
    if not workbook.exists():
        candidates = list(raw_dir.glob("*.xlsx"))
        workbook = next(
            (
                path
                for path in candidates
                if "cross" in path.name.lower() or "soc" in path.name.lower()
            ),
            workbook,
        )
    if not workbook.exists():
        return {}
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RuntimeError(
            "Crosswalk parsing requires optional dependency: pip install -e .[excel]"
        ) from error
    workbook_handle = load_workbook(workbook, read_only=True, data_only=True)
    sheet = workbook_handle.active
    rows = list(sheet.iter_rows(values_only=True))
    workbook_handle.close()
    header_index = next(
        (
            index
            for index, row in enumerate(rows)
            if any("2019" in str(cell) and "Code" in str(cell) for cell in row)
        ),
        0,
    )
    header = [str(cell or "").strip() for cell in rows[header_index]]
    source_index = next(
        (
            index
            for index, item in enumerate(header)
            if "2019" in item and "code" in item.lower()
        ),
        0,
    )
    target_index = next(
        (
            index
            for index, item in enumerate(header)
            if "2018" in item and "code" in item.lower()
        ),
        1,
    )
    mapping: dict[str, list[str]] = {}
    for row in rows[header_index + 1 :]:
        if len(row) <= max(source_index, target_index):
            continue
        source = str(row[source_index] or "").strip()
        target = str(row[target_index] or "").strip()
        if source and target:
            mapping.setdefault(source, []).append(target)
    return mapping


def _load_aioe(raw_dir: Path) -> dict[str, float]:
    candidates = [raw_dir / "aioe.csv", raw_dir / "AIOE.csv"]
    path = next((candidate for candidate in candidates if candidate.exists()), None)
    if path:
        rows = read_csv(path)
    else:
        workbook = raw_dir / "aioe_data_appendix.xlsx"
        if not workbook.exists():
            return {}
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError(
                "AIOE workbook parsing requires optional dependency: pip install -e .[excel]"
            ) from error
        workbook_handle = load_workbook(workbook, read_only=True, data_only=True)
        sheet = workbook_handle["Appendix A"]
        values = list(sheet.iter_rows(values_only=True))
        workbook_handle.close()
        rows = [dict(zip(values[0], row)) for row in values[1:]]
    result = {}
    for row in rows:
        code = _value(row, "soc", "soc_code", "soc 2018 code", "occupation code")
        score = _number(
            _value(row, "aioe", "ai exposure", "exposure"), non_negative=True
        )
        if code and score is not None:
            result[code] = score
    return result


def _read_csv_or_zip(path: Path) -> list[dict[str, str]]:
    if path.suffix.lower() != ".zip":
        return read_csv(path)
    with zipfile.ZipFile(path) as archive:
        spreadsheet_names = [
            name for name in archive.namelist() if name.lower().endswith(".xlsx")
        ]
        if spreadsheet_names:
            try:
                from openpyxl import load_workbook
            except ImportError as error:
                raise RuntimeError(
                    "XLSX parsing requires optional dependency: pip install -e .[excel]"
                ) from error
            workbook_handle = load_workbook(
                io.BytesIO(archive.read(spreadsheet_names[0])),
                read_only=True,
                data_only=True,
            )
            sheet = workbook_handle.active
            values = list(sheet.iter_rows(values_only=True))
            workbook_handle.close()
            if not values:
                return []
            header = [str(cell or "").strip() for cell in values[0]]
            return [
                dict(zip(header, row))
                for row in values[1:]
                if any(cell not in (None, "") for cell in row)
            ]
        names = [
            name
            for name in archive.namelist()
            if name.lower().endswith((".csv", ".txt"))
        ]
        if not names:
            return []
        with archive.open(names[0]) as raw:
            content = io.TextIOWrapper(raw, encoding="utf-8-sig", errors="replace")
            sample = content.read(2048)
            content.seek(0)
            delimiter = "\t" if "\t" in sample else ","
            return list(csv.DictReader(content, delimiter=delimiter))


def _load_oews(raw_dir: Path) -> dict[str, dict[str, float | None]]:
    candidates = [raw_dir / "bls_oews_may_2025.zip", raw_dir / "bls_oews.csv"]
    path = next((candidate for candidate in candidates if candidate.exists()), None)
    if not path:
        return {}
    result = {}
    for row in _read_csv_or_zip(path):
        code = _value(row, "occ code", "occ_code", "soc code", "occupation code")
        if not code or code.lower() in {"code", "occ_code"}:
            continue
        result[code] = {
            "employment_2024": _number(
                _value(row, "tot emp", "total employment", "employment"),
                non_negative=True,
            ),
            "median_annual_wage": _number(
                _value(row, "a median", "median annual wage", "annual median wage"),
                non_negative=True,
            ),
        }
    return result


def _load_projections(raw_dir: Path) -> dict[str, dict[str, float | None]]:
    candidates = [
        raw_dir / "bls_projections_2024_2034.xlsx",
        raw_dir / "bls_projections_2024_2034.csv",
    ]
    path = next((candidate for candidate in candidates if candidate.exists()), None)
    if not path:
        return {}
    if path.suffix.lower() == ".csv":
        rows = read_csv(path)
    else:
        try:
            from openpyxl import load_workbook
        except ImportError as error:
            raise RuntimeError(
                "Projection parsing requires optional dependency: pip install -e .[excel]"
            ) from error
        workbook_handle = load_workbook(path, read_only=True, data_only=True)
        sheet = (
            workbook_handle["Table 1.2"]
            if "Table 1.2" in workbook_handle.sheetnames
            else workbook_handle.active
        )
        values = list(sheet.iter_rows(values_only=True))
        workbook_handle.close()
        header_index = next(
            (
                index
                for index, row in enumerate(values)
                if any(
                    "2024 national employment matrix code" in str(cell).lower()
                    for cell in row
                )
                and any(str(cell).strip().lower() == "employment, 2024" for cell in row)
            ),
            0,
        )
        header = [str(cell or "").strip() for cell in values[header_index]]
        rows = [dict(zip(header, row)) for row in values[header_index + 1 :]]
    result = {}
    for row in rows:
        annual_openings_thousands = _number(
            _value(
                row,
                "Occupational openings, 2024-34 annual average",
                "occupational openings annual average",
            ),
            non_negative=True,
        )
        code = _value(
            row,
            "2024 National Employment Matrix code",
            "national employment matrix code",
            "occupation code",
            "soc code",
        )
        if not code or code.lower() in {"code", "occupation code"}:
            continue
        result[code] = {
            "projected_employment_2024_thousands": _number(
                _value(row, "Employment, 2024", "employment 2024"),
                non_negative=True,
            ),
            "projected_employment_2034_thousands": _number(
                _value(row, "Employment, 2034", "employment 2034"),
                non_negative=True,
            ),
            "employment_change_2024_2034_pct": _number(
                _value(
                    row,
                    "Employment change, percent, 2024-34",
                    "employment change percent 2024 34",
                )
            ),
            "annual_openings_2024_2034": (
                annual_openings_thousands * 1000
                if annual_openings_thousands is not None
                else None
            ),
            "projection_median_annual_wage": _number(
                _value(
                    row,
                    "Median annual wage, dollars, 2024",
                    "median annual wage dollars 2024",
                ),
                non_negative=True,
            ),
        }
    return result


def build_dataset(
    raw_dir: Path, processed_dir: Path, demo: bool = False
) -> dict[str, object]:
    processed_dir.mkdir(parents=True, exist_ok=True)
    if demo:
        rows = demo_rows()
        tasks = demo_tasks()
        source_flags = {"mode": "demo", "rows": len(rows)}
    else:
        onet = _load_onet(raw_dir)
        tasks = _load_tasks(raw_dir)
        occupation_codes = {str(item["onet_soc_code"]) for item in onet}
        for task in tasks:
            if task["onet_soc_code"] not in occupation_codes:
                task["task_quality_flags"] = "task_without_occupation_record"
        crosswalk = _load_crosswalk(raw_dir)
        soc_to_onet: dict[str, set[str]] = {}
        for source_code, targets in crosswalk.items():
            for target_code in set(targets):
                soc_to_onet.setdefault(target_code, set()).add(source_code)
        aioe = _load_aioe(raw_dir)
        oews = _load_oews(raw_dir)
        projections = _load_projections(raw_dir)
        rows = []
        for item in onet:
            codes = [
                code
                for code in list(
                    dict.fromkeys(crosswalk.get(str(item["onet_soc_code"]), [""]))
                )
                if str(code).strip()
            ]
            codes = codes or [""]
            crosswalk_weight = 1.0 / len([code for code in codes if code]) if any(codes) else 1.0
            for soc_code in codes or [""]:
                wage = oews.get(soc_code, {})
                projection = projections.get(soc_code, {})
                exposure = aioe.get(soc_code)
                flags = []
                if len([code for code in codes if code]) > 1:
                    flags.append("multiple_soc_crosswalk")
                    flags.append("uniform_crosswalk_fallback")
                if not soc_code:
                    flags.append("missing_crosswalk")
                if soc_code and len(soc_to_onet.get(soc_code, set())) > 1:
                    flags.append("shared_soc_crosswalk")
                if exposure is None:
                    flags.append("missing_aioe")
                if not wage:
                    flags.append("missing_oews")
                if not projection:
                    flags.append("missing_projections")
                rows.append(
                    {
                        "onet_soc_code": item["onet_soc_code"],
                        "soc_2018_code": soc_code,
                        "crosswalk_weight": crosswalk_weight,
                        "title": item["title"],
                        "description": item["description"],
                        "ai_exposure": exposure,
                        "ai_exposure_source": "AIOE" if exposure is not None else "",
                        "employment_2024": wage.get("employment_2024"),
                        "projected_employment_2024_thousands": projection.get(
                            "projected_employment_2024_thousands"
                        ),
                        "projected_employment_2034_thousands": projection.get(
                            "projected_employment_2034_thousands"
                        ),
                        "employment_change_2024_2034_pct": projection.get(
                            "employment_change_2024_2034_pct"
                        ),
                        "annual_openings_2024_2034": projection.get(
                            "annual_openings_2024_2034"
                        ),
                        "median_annual_wage": wage.get("median_annual_wage")
                        or projection.get("projection_median_annual_wage"),
                        "onet_version": "30.3",
                        "wage_vintage": "May 2025",
                        "projection_vintage": "2024-2034",
                        "crosswalk_method": "official_onet_crosswalk",
                        "data_quality_flags": ";".join(flags),
                    }
                )
        source_flags = {
            "mode": "public_raw",
            "rows": len(rows),
            "onet_rows": len(onet),
            "task_rows": len(tasks),
            "task_occupations": len({task["onet_soc_code"] for task in tasks}),
            "task_unknown_occupation_rows": sum(
                task["task_quality_flags"] == "task_without_occupation_record"
                for task in tasks
            ),
            "crosswalk_sources": len(crosswalk),
            "aioe_rows": len(aioe),
            "oews_rows": len(oews),
            "projection_rows": len(projections),
        }
    if not demo:
        validate_alias_registry(load_alias_registry(), rows)
        missing_layers = [
            label
            for label, values in (
                ("O*NET occupations", onet),
                ("O*NET tasks", tasks),
                ("O*NET-to-SOC crosswalk", crosswalk),
                ("AIOE exposure", aioe),
                ("BLS OEWS", oews),
                ("BLS Employment Projections", projections),
            )
            if not values
        ]
        if missing_layers:
            raise ValueError(
                "full build requires every registered data layer; missing or "
                "unreadable: "
                + ", ".join(missing_layers)
                + ". Run 'atlas download' and resolve every failed source first."
            )
    count = write_csv(processed_dir / "occupations.csv", rows, FIELDS)
    task_count = write_csv(processed_dir / "tasks.csv", tasks, TASK_FIELDS)
    manifest = {
        "schema": FIELDS,
        "row_count": count,
        "task_schema": TASK_FIELDS,
        "task_row_count": task_count,
        "build": source_flags,
        "sha256": sha256(processed_dir / "occupations.csv"),
        "tasks_sha256": sha256(processed_dir / "tasks.csv"),
    }
    write_json(processed_dir / "data_manifest.json", manifest)
    return manifest
